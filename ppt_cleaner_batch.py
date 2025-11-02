#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PowerPoint Cleaner Tool - Batch Processing
Cong cu don dep PowerPoint: Xu ly NHIEU file cung luc
Khong can mo PowerPoint, chay doc lap
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from pptx import Presentation
from pptx.util import Inches
import os
import shutil
from datetime import datetime
from pathlib import Path
import re
from threading import Thread
import queue

class PowerPointCleanerBatch:
    def __init__(self, root):
        self.root = root
        self.root.title("PowerPoint Cleaner - Xu ly nhieu file")
        self.root.geometry("950x850")
        self.root.resizable(True, True)  # Cho phep dieu chinh kich thuoc
        self.root.minsize(800, 700)  # Kich thuoc toi thieu
        
        # Variables
        self.file_paths = []  # Danh sach files
        self.auto_backup = tk.BooleanVar(value=True)
        self.processing = False
        self.result_queue = queue.Queue()
        
        self.create_widgets()
        
    def create_widgets(self):
        """Tao giao dien"""
        
        # Header
        header_frame = tk.Frame(self.root, bg="#2c3e50", height=80)
        header_frame.pack(fill=tk.X)
        
        title_label = tk.Label(
            header_frame,
            text="PowerPoint Cleaner - Batch",
            font=("Arial", 20, "bold"),
            bg="#2c3e50",
            fg="white"
        )
        title_label.pack(pady=5)
        
        subtitle_label = tk.Label(
            header_frame,
            text="Xu ly NHIEU file cung luc - Tiet kiem thoi gian",
            font=("Arial", 10),
            bg="#2c3e50",
            fg="#ecf0f1"
        )
        subtitle_label.pack()
        
        # Main content
        main_frame = tk.Frame(self.root, padx=20, pady=15)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # File selection
        file_frame = tk.LabelFrame(main_frame, text="1. Chon cac file PowerPoint", font=("Arial", 11, "bold"), padx=10, pady=10)
        file_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        
        # Buttons
        btn_frame = tk.Frame(file_frame)
        btn_frame.pack(fill=tk.X, pady=(0, 10))
        
        add_files_btn = tk.Button(
            btn_frame,
            text="+ Them files...",
            command=self.add_files,
            bg="#3498db",
            fg="white",
            font=("Arial", 10),
            cursor="hand2",
            relief=tk.FLAT,
            padx=15
        )
        add_files_btn.pack(side=tk.LEFT, padx=(0, 5))
        
        add_folder_btn = tk.Button(
            btn_frame,
            text="+ Them folder...",
            command=self.add_folder,
            bg="#3498db",
            fg="white",
            font=("Arial", 10),
            cursor="hand2",
            relief=tk.FLAT,
            padx=15
        )
        add_folder_btn.pack(side=tk.LEFT, padx=5)
        
        clear_btn = tk.Button(
            btn_frame,
            text="Xoa tat ca",
            command=self.clear_files,
            bg="#e74c3c",
            fg="white",
            font=("Arial", 10),
            cursor="hand2",
            relief=tk.FLAT,
            padx=15
        )
        clear_btn.pack(side=tk.RIGHT)
        
        # File list
        list_frame = tk.Frame(file_frame)
        list_frame.pack(fill=tk.BOTH, expand=True)
        
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.file_listbox = tk.Listbox(
            list_frame,
            font=("Courier", 9),
            yscrollcommand=scrollbar.set,
            selectmode=tk.EXTENDED
        )
        self.file_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=self.file_listbox.yview)
        
        # Right-click menu
        self.context_menu = tk.Menu(self.file_listbox, tearoff=0)
        self.context_menu.add_command(label="Xoa file da chon", command=self.remove_selected)
        self.file_listbox.bind("<Button-3>", self.show_context_menu)
        
        self.file_count_label = tk.Label(file_frame, text="Chua chon file nao", font=("Arial", 9), fg="gray")
        self.file_count_label.pack(pady=(5, 0))
        
        # Notebook (Tabs)
        self.notebook = ttk.Notebook(main_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
        
        # Tab 1: Xoa hinh anh
        self.create_image_tab()
        
        # Tab 2: Loc text
        self.create_text_tab()
        
        # Options
        options_frame = tk.LabelFrame(main_frame, text="Tuy chon", font=("Arial", 11, "bold"), padx=10, pady=10)
        options_frame.pack(fill=tk.X, pady=(0, 10))
        
        backup_cb = tk.Checkbutton(
            options_frame,
            text="Tu dong backup moi file truoc khi xu ly",
            variable=self.auto_backup,
            font=("Arial", 10)
        )
        backup_cb.pack(anchor="w")
        
        # Progress
        self.progress_frame = tk.LabelFrame(main_frame, text="Tien trinh", font=("Arial", 11, "bold"), padx=10, pady=10)
        self.progress_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.progress_bar = ttk.Progressbar(self.progress_frame, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=(0, 5))
        
        self.progress_label = tk.Label(self.progress_frame, text="Chua bat dau", font=("Arial", 9))
        self.progress_label.pack()
        
        # Action buttons
        button_frame = tk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=(10, 0))
        
        self.process_btn = tk.Button(
            button_frame,
            text="? XU LY TAT CA FILES",
            command=self.process_all_files,
            bg="#27ae60",
            fg="white",
            font=("Arial", 13, "bold"),
            cursor="hand2",
            relief=tk.RAISED,
            padx=40,
            pady=15,
            borderwidth=3
        )
        self.process_btn.pack(fill=tk.X, ipady=5)
        
        # Status bar
        self.status_label = tk.Label(
            self.root,
            text="San sang - Them files de bat dau",
            font=("Arial", 9),
            bg="#ecf0f1",
            fg="#2c3e50",
            anchor="w",
            padx=10,
            pady=5
        )
        self.status_label.pack(side=tk.BOTTOM, fill=tk.X)
    
    def create_image_tab(self):
        """Tab xoa hinh anh"""
        image_frame = ttk.Frame(self.notebook)
        self.notebook.add(image_frame, text=" Xoa hinh anh")
        
        # Variables
        self.img_width = tk.DoubleVar(value=1.6)
        self.img_height = tk.DoubleVar(value=1.6)
        self.img_tolerance = tk.DoubleVar(value=0.01)
        self.img_mode = tk.StringVar(value="and")
        
        # Size settings
        size_frame = tk.LabelFrame(image_frame, text="Kich thuoc (inch)", font=("Arial", 10, "bold"), padx=10, pady=10)
        size_frame.pack(fill=tk.X, padx=20, pady=10)
        
        # Width
        width_frame = tk.Frame(size_frame)
        width_frame.pack(fill=tk.X, pady=3)
        tk.Label(width_frame, text="Chieu rong:", font=("Arial", 9), width=15, anchor="w").pack(side=tk.LEFT)
        tk.Spinbox(width_frame, from_=0.1, to=20.0, increment=0.1, textvariable=self.img_width, 
                   font=("Arial", 9), width=8).pack(side=tk.LEFT, padx=5)
        tk.Label(width_frame, text="inch", font=("Arial", 9)).pack(side=tk.LEFT)
        
        # Height
        height_frame = tk.Frame(size_frame)
        height_frame.pack(fill=tk.X, pady=3)
        tk.Label(height_frame, text="Chieu cao:", font=("Arial", 9), width=15, anchor="w").pack(side=tk.LEFT)
        tk.Spinbox(height_frame, from_=0.1, to=20.0, increment=0.1, textvariable=self.img_height, 
                   font=("Arial", 9), width=8).pack(side=tk.LEFT, padx=5)
        tk.Label(height_frame, text="inch", font=("Arial", 9)).pack(side=tk.LEFT)
        
        # Tolerance
        tol_frame = tk.Frame(size_frame)
        tol_frame.pack(fill=tk.X, pady=3)
        tk.Label(tol_frame, text="Dung sai:", font=("Arial", 9), width=15, anchor="w").pack(side=tk.LEFT)
        tk.Spinbox(tol_frame, from_=0.001, to=0.5, increment=0.01, textvariable=self.img_tolerance, 
                   font=("Arial", 9), width=8).pack(side=tk.LEFT, padx=5)
        tk.Label(tol_frame, text="inch", font=("Arial", 9)).pack(side=tk.LEFT)
        
        # Match mode
        mode_frame = tk.LabelFrame(image_frame, text="Che do so khop", font=("Arial", 10, "bold"), padx=10, pady=5)
        mode_frame.pack(fill=tk.X, padx=20)
        
        modes = [
            ("CA width VA height khop (AND)", "and"),
            ("width HOAC height khop (OR)", "or"),
            ("Chi theo width", "width_only"),
            ("Chi theo height", "height_only")
        ]
        
        for text, value in modes:
            tk.Radiobutton(mode_frame, text=text, variable=self.img_mode, value=value, 
                          font=("Arial", 9)).pack(anchor="w", pady=1)
    
    def create_text_tab(self):
        """Tab loc text"""
        text_frame = ttk.Frame(self.notebook)
        self.notebook.add(text_frame, text=" Loc text")
        
        # Variables
        self.text_mode = tk.StringVar(value="keep_vietnamese")
        self.delete_empty_shapes = tk.BooleanVar(value=True)
        
        # Mode selection
        mode_frame = tk.LabelFrame(text_frame, text="Che do loc", font=("Arial", 10, "bold"), padx=10, pady=10)
        mode_frame.pack(fill=tk.X, padx=20, pady=10)
        
        modes = [
            ("Giu tieng Viet - Xoa tieng Anh", "keep_vietnamese"),
            ("Chi giu tieng Anh - Xoa tieng Viet", "keep_english"),
            ("Xoa tat ca text", "delete_all")
        ]
        
        for text, value in modes:
            tk.Radiobutton(mode_frame, text=text, variable=self.text_mode, value=value,
                          font=("Arial", 9)).pack(anchor="w", pady=2)
        
        # Options
        opt_frame = tk.LabelFrame(text_frame, text="Tuy chon", font=("Arial", 10, "bold"), padx=10, pady=10)
        opt_frame.pack(fill=tk.X, padx=20)
        
        tk.Checkbutton(opt_frame, text="Xoa textbox rong sau khi loc", 
                      variable=self.delete_empty_shapes, font=("Arial", 9)).pack(anchor="w")
    
    def add_files(self):
        """Them files"""
        filenames = filedialog.askopenfilenames(
            title="Chon cac file PowerPoint",
            filetypes=[("PowerPoint files", "*.pptx"), ("All files", "*.*")]
        )
        if filenames:
            for filename in filenames:
                if filename not in self.file_paths:
                    self.file_paths.append(filename)
                    self.file_listbox.insert(tk.END, os.path.basename(filename))
            self.update_file_count()
    
    def add_folder(self):
        """Them tat ca file .pptx trong folder"""
        folder = filedialog.askdirectory(title="Chon folder chua cac file PowerPoint")
        if folder:
            added = 0
            for file in os.listdir(folder):
                if file.endswith('.pptx') and not file.startswith('~'):
                    full_path = os.path.join(folder, file)
                    if full_path not in self.file_paths:
                        self.file_paths.append(full_path)
                        self.file_listbox.insert(tk.END, file)
                        added += 1
            
            if added > 0:
                self.update_file_count()
                messagebox.showinfo("Thanh cong", f"Da them {added} file tu folder!")
            else:
                messagebox.showinfo("Thong bao", "Khong tim thay file .pptx nao trong folder!")
    
    def clear_files(self):
        """Xoa tat ca files"""
        if self.file_paths and messagebox.askyesno("Xac nhan", "Xoa tat ca files da chon?"):
            self.file_paths.clear()
            self.file_listbox.delete(0, tk.END)
            self.update_file_count()
    
    def remove_selected(self):
        """Xoa files da chon"""
        selected = self.file_listbox.curselection()
        if selected:
            for index in reversed(selected):
                self.file_listbox.delete(index)
                del self.file_paths[index]
            self.update_file_count()
    
    def show_context_menu(self, event):
        """Hien context menu"""
        if self.file_listbox.curselection():
            self.context_menu.post(event.x_root, event.y_root)
    
    def update_file_count(self):
        """Cap nhat so luong files"""
        count = len(self.file_paths)
        if count == 0:
            self.file_count_label.config(text="Chua chon file nao", fg="gray")
            self.status_label.config(text="San sang - Them files de bat dau")
        else:
            self.file_count_label.config(text=f"Da chon: {count} file", fg="green")
            self.status_label.config(text=f"San sang xu ly {count} file")
    
    def process_all_files(self):
        """Xu ly tat ca files"""
        if not self.file_paths:
            messagebox.showwarning("Canh bao", "Vui long chon it nhat 1 file!")
            return
        
        if self.processing:
            messagebox.showwarning("Canh bao", "Dang xu ly, vui long doi!")
            return
        
        # Confirm
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        action = "xoa hinh anh" if "hinh anh" in current_tab.lower() else "loc text"
        
        if not messagebox.askyesno("Xac nhan", 
            f"Xu ly {len(self.file_paths)} file voi tac vu: {action}?\n\n"
            "Thao tac nay khong the hoan tac!"):
            return
        
        # Start processing in thread
        self.processing = True
        self.process_btn.config(state=tk.DISABLED, text="Dang xu ly...")
        
        thread = Thread(target=self.process_files_thread, daemon=True)
        thread.start()
        
        # Monitor progress
        self.root.after(100, self.check_progress)
    
    def process_files_thread(self):
        """Xu ly files trong thread rieng"""
        total = len(self.file_paths)
        results = []
        
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        is_image_mode = "hinh anh" in current_tab.lower()
        
        for idx, file_path in enumerate(self.file_paths, 1):
            try:
                # Update progress
                progress = (idx / total) * 100
                filename = os.path.basename(file_path)
                self.result_queue.put(('progress', progress, f"Dang xu ly ({idx}/{total}): {filename}"))
                
                # Backup
                if self.auto_backup.get():
                    self.create_backup(file_path)
                
                # Process
                if is_image_mode:
                    result = self.process_single_image_file(file_path)
                else:
                    result = self.process_single_text_file(file_path)
                
                results.append({'file': filename, 'status': 'success', 'data': result})
                
            except Exception as e:
                results.append({'file': os.path.basename(file_path), 'status': 'error', 'data': str(e)})
        
        # Done
        self.result_queue.put(('done', results))
    
    def check_progress(self):
        """Kiem tra tien trinh"""
        try:
            while True:
                msg = self.result_queue.get_nowait()
                
                if msg[0] == 'progress':
                    _, progress, text = msg
                    self.progress_bar['value'] = progress
                    self.progress_label.config(text=text)
                    self.status_label.config(text=text)
                    
                elif msg[0] == 'done':
                    _, results = msg
                    self.processing = False
                    self.process_btn.config(state=tk.NORMAL, text="? XU LY TAT CA FILES")
                    self.progress_bar['value'] = 100
                    self.progress_label.config(text="Hoan tat!")
                    self.show_results(results)
                    return
        except queue.Empty:
            pass
        
        if self.processing:
            self.root.after(100, self.check_progress)
    
    def process_single_image_file(self, file_path):
        """Xu ly xoa hinh anh cho 1 file"""
        prs = Presentation(file_path)
        deleted_count = 0
        total_images = 0
        
        for slide in prs.slides:
            shapes_to_delete = []
            
            for shape in slide.shapes:
                if shape.shape_type == 13:
                    total_images += 1
                    width_inches = shape.width / 914400
                    height_inches = shape.height / 914400
                    
                    if self.check_image_match(width_inches, height_inches):
                        shapes_to_delete.append(shape)
            
            for shape in shapes_to_delete:
                sp = shape.element
                sp.getparent().remove(sp)
                deleted_count += 1
        
        prs.save(file_path)
        return {'total': total_images, 'deleted': deleted_count}
    
    def process_single_text_file(self, file_path):
        """Xu ly loc text cho 1 file"""
        prs = Presentation(file_path)
        modified_count = 0
        deleted_count = 0
        total_textboxes = 0
        
        for slide in prs.slides:
            shapes_to_delete = []
            
            for shape in slide.shapes:
                if self.has_text(shape):
                    total_textboxes += 1
                    original_text = self.get_shape_text(shape)
                    new_text = self.filter_text(original_text)
                    
                    if new_text != original_text:
                        if len(new_text.strip()) == 0 and self.delete_empty_shapes.get():
                            shapes_to_delete.append(shape)
                            deleted_count += 1
                        else:
                            self.set_shape_text(shape, new_text)
                            modified_count += 1
            
            for shape in shapes_to_delete:
                sp = shape.element
                sp.getparent().remove(sp)
        
        prs.save(file_path)
        return {'total': total_textboxes, 'modified': modified_count, 'deleted': deleted_count}
    
    def show_results(self, results):
        """Hien thi ket qua"""
        result_window = tk.Toplevel(self.root)
        result_window.title("Ket qua xu ly")
        result_window.geometry("700x500")
        
        # Header
        success_count = sum(1 for r in results if r['status'] == 'success')
        error_count = len(results) - success_count
        
        header = tk.Label(
            result_window,
            text=f"Hoan tat! Thanh cong: {success_count} | Loi: {error_count}",
            font=("Arial", 12, "bold"),
            bg="#27ae60" if error_count == 0 else "#f39c12",
            fg="white",
            pady=10
        )
        header.pack(fill=tk.X)
        
        # Content
        content_frame = tk.Frame(result_window, padx=10, pady=10)
        content_frame.pack(fill=tk.BOTH, expand=True)
        
        text_widget = scrolledtext.ScrolledText(content_frame, font=("Courier", 9), wrap=tk.WORD)
        text_widget.pack(fill=tk.BOTH, expand=True)
        
        current_tab = self.notebook.tab(self.notebook.select(), "text")
        is_image_mode = "hinh anh" in current_tab.lower()
        
        for idx, result in enumerate(results, 1):
            text_widget.insert(tk.END, f"\n{'='*70}\n")
            text_widget.insert(tk.END, f"{idx}. {result['file']}\n")
            text_widget.insert(tk.END, f"{'='*70}\n")
            
            if result['status'] == 'success':
                text_widget.insert(tk.END, "Trang thai: THANH CONG\n")
                data = result['data']
                
                if is_image_mode:
                    text_widget.insert(tk.END, f"Tong hinh anh: {data['total']}\n")
                    text_widget.insert(tk.END, f"Da xoa: {data['deleted']}\n")
                else:
                    text_widget.insert(tk.END, f"Tong textbox: {data['total']}\n")
                    text_widget.insert(tk.END, f"Da cap nhat: {data['modified']}\n")
                    text_widget.insert(tk.END, f"Da xoa: {data['deleted']}\n")
            else:
                text_widget.insert(tk.END, f"Trang thai: LOI\n")
                text_widget.insert(tk.END, f"Chi tiet: {result['data']}\n")
        
        text_widget.config(state=tk.DISABLED)
        
        # Close button
        close_btn = tk.Button(
            result_window,
            text="Dong",
            command=result_window.destroy,
            bg="#95a5a6",
            fg="white",
            font=("Arial", 10),
            padx=30,
            pady=5
        )
        close_btn.pack(pady=10)
    
    # Helper methods
    def check_image_match(self, width_inches, height_inches):
        target_width = self.img_width.get()
        target_height = self.img_height.get()
        tolerance = self.img_tolerance.get()
        mode = self.img_mode.get()
        
        width_match = abs(width_inches - target_width) < tolerance
        height_match = abs(height_inches - target_height) < tolerance
        
        if mode == "and":
            return width_match and height_match
        elif mode == "or":
            return width_match or height_match
        elif mode == "width_only":
            return width_match
        elif mode == "height_only":
            return height_match
        return False
    
    def has_text(self, shape):
        try:
            return shape.has_text_frame and shape.text_frame.text.strip()
        except:
            return False
    
    def get_shape_text(self, shape):
        try:
            return shape.text_frame.text
        except:
            return ""
    
    def set_shape_text(self, shape, text):
        try:
            shape.text_frame.text = text
        except:
            pass
    
    def filter_text(self, text):
        mode = self.text_mode.get()
        
        if mode == "keep_vietnamese":
            return self.keep_unicode_lines(text)
        elif mode == "keep_english":
            return self.keep_ascii_lines(text)
        elif mode == "delete_all":
            return ""
        return text
    
    def keep_unicode_lines(self, text):
        lines = text.split('\n')
        result = [line for line in lines if self.has_unicode_char(line)]
        return '\n'.join(result)
    
    def keep_ascii_lines(self, text):
        lines = text.split('\n')
        result = [line for line in lines if not self.has_unicode_char(line) and line.strip()]
        return '\n'.join(result)
    
    def has_unicode_char(self, line):
        if not line.strip():
            return False
        return any(ord(char) > 127 for char in line)
    
    def create_backup(self, file_path):
        file_dir = os.path.dirname(file_path)
        file_name = os.path.basename(file_path)
        name_without_ext = os.path.splitext(file_name)[0]
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"{name_without_ext}_backup_{timestamp}.pptx"
        backup_path = os.path.join(file_dir, backup_name)
        
        shutil.copy2(file_path, backup_path)
        return backup_path


def main():
    root = tk.Tk()
    app = PowerPointCleanerBatch(root)
    root.mainloop()


if __name__ == "__main__":
    main()
