#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PowerPoint Image Remover Tool
C?ng c? x?a h?nh ?nh theo k?ch th??c trong file PowerPoint
Kh?ng c?n m? PowerPoint, ch?y ??c l?p
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pptx import Presentation
from pptx.util import Inches
import os
import shutil
from datetime import datetime
from pathlib import Path

class PowerPointImageRemover:
    def __init__(self, root):
        self.root = root
        self.root.title("PowerPoint Image Remover - C?ng c? x?a h?nh ?nh")
        self.root.geometry("700x600")
        self.root.resizable(False, False)
        
        self.root.minsize(700, 600)  # Kich thuoc toi thieu
        # Variables
        self.file_path = tk.StringVar()
        self.width_value = tk.DoubleVar(value=1.6)
        self.height_value = tk.DoubleVar(value=1.6)
        self.match_mode = tk.StringVar(value="and")
        self.tolerance = tk.DoubleVar(value=0.01)
        self.auto_backup = tk.BooleanVar(value=True)
        
        self.create_widgets()
        
    def create_widgets(self):
        """T?o giao di?n"""
        
        # Header
        header_frame = tk.Frame(self.root, bg="#2c3e50", height=80)
        header_frame.pack(fill=tk.X)
        
        title_label = tk.Label(
            header_frame,
            text="PowerPoint Image Remover",
            font=("Arial", 20, "bold"),
            bg="#2c3e50",
            fg="white"
        )
        title_label.pack(pady=10)
        
        subtitle_label = tk.Label(
            header_frame,
            text="X?a h?nh ?nh theo k?ch th??c - Kh?ng c?n m? PowerPoint",
            font=("Arial", 10),
            bg="#2c3e50",
            fg="#ecf0f1"
        )
        subtitle_label.pack()
        
        # Main content
        main_frame = tk.Frame(self.root, padx=20, pady=20)
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # File selection
        file_frame = tk.LabelFrame(main_frame, text="1. Ch?n file PowerPoint", font=("Arial", 11, "bold"), padx=10, pady=10)
        file_frame.pack(fill=tk.X, pady=(0, 15))
        
        file_entry = tk.Entry(file_frame, textvariable=self.file_path, font=("Arial", 10), state="readonly")
        file_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 10))
        
        browse_btn = tk.Button(
            file_frame,
            text="Ch?n file...",
            command=self.browse_file,
            bg="#3498db",
            fg="white",
            font=("Arial", 10),
            cursor="hand2",
            relief=tk.FLAT,
            padx=20
        )
        browse_btn.pack(side=tk.RIGHT)
        
        # Size settings
        size_frame = tk.LabelFrame(main_frame, text="2. C?i ??t k?ch th??c (inch)", font=("Arial", 11, "bold"), padx=10, pady=10)
        size_frame.pack(fill=tk.X, pady=(0, 15))
        
        # Width
        width_frame = tk.Frame(size_frame)
        width_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(width_frame, text="Chi?u r?ng (Width):", font=("Arial", 10), width=20, anchor="w").pack(side=tk.LEFT)
        width_spinbox = tk.Spinbox(
            width_frame,
            from_=0.1,
            to=20.0,
            increment=0.1,
            textvariable=self.width_value,
            font=("Arial", 10),
            width=10
        )
        width_spinbox.pack(side=tk.LEFT, padx=10)
        tk.Label(width_frame, text="inch", font=("Arial", 10)).pack(side=tk.LEFT)
        
        # Height
        height_frame = tk.Frame(size_frame)
        height_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(height_frame, text="Chi?u cao (Height):", font=("Arial", 10), width=20, anchor="w").pack(side=tk.LEFT)
        height_spinbox = tk.Spinbox(
            height_frame,
            from_=0.1,
            to=20.0,
            increment=0.1,
            textvariable=self.height_value,
            font=("Arial", 10),
            width=10
        )
        height_spinbox.pack(side=tk.LEFT, padx=10)
        tk.Label(height_frame, text="inch", font=("Arial", 10)).pack(side=tk.LEFT)
        
        # Tolerance
        tolerance_frame = tk.Frame(size_frame)
        tolerance_frame.pack(fill=tk.X, pady=5)
        
        tk.Label(tolerance_frame, text="Dung sai (Tolerance):", font=("Arial", 10), width=20, anchor="w").pack(side=tk.LEFT)
        tolerance_spinbox = tk.Spinbox(
            tolerance_frame,
            from_=0.001,
            to=0.5,
            increment=0.01,
            textvariable=self.tolerance,
            font=("Arial", 10),
            width=10
        )
        tolerance_spinbox.pack(side=tk.LEFT, padx=10)
        tk.Label(tolerance_frame, text="inch", font=("Arial", 10)).pack(side=tk.LEFT)
        
        # Match mode
        mode_frame = tk.LabelFrame(main_frame, text="3. Ch? ?? so kh?p", font=("Arial", 11, "bold"), padx=10, pady=10)
        mode_frame.pack(fill=tk.X, pady=(0, 15))
        
        modes = [
            ("Ch? x?a khi C? width V? height kh?p (AND)", "and"),
            ("X?a khi width HO?C height kh?p (OR)", "or"),
            ("Ch? x?a theo width (b?t k? height)", "width_only"),
            ("Ch? x?a theo height (b?t k? width)", "height_only")
        ]
        
        for text, value in modes:
            rb = tk.Radiobutton(
                mode_frame,
                text=text,
                variable=self.match_mode,
                value=value,
                font=("Arial", 10),
                anchor="w"
            )
            rb.pack(fill=tk.X, pady=2)
        
        # Options
        options_frame = tk.LabelFrame(main_frame, text="4. T?y ch?n", font=("Arial", 11, "bold"), padx=10, pady=10)
        options_frame.pack(fill=tk.X, pady=(0, 15))
        
        backup_cb = tk.Checkbutton(
            options_frame,
            text="T? ??ng backup file g?c tr??c khi x?a (khuy?n ngh?)",
            variable=self.auto_backup,
            font=("Arial", 10)
        )
        backup_cb.pack(anchor="w")
        
        # Action buttons
        button_frame = tk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=(10, 0))
        
        preview_btn = tk.Button(
            button_frame,
            text="Preview - Xem tr??c",
            command=self.preview_images,
            bg="#f39c12",
            fg="white",
            font=("Arial", 11, "bold"),
            cursor="hand2",
            relief=tk.FLAT,
            padx=30,
            pady=10
        )
        preview_btn.pack(side=tk.LEFT, padx=(0, 10), fill=tk.X, expand=True)
        
        process_btn = tk.Button(
            button_frame,
            text="X?a h?nh ?nh",
            command=self.process_file,
            bg="#27ae60",
            fg="white",
            font=("Arial", 11, "bold"),
            cursor="hand2",
            relief=tk.FLAT,
            padx=30,
            pady=10
        )
        process_btn.pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        # Status bar
        self.status_label = tk.Label(
            self.root,
            text="S?n s?ng",
            font=("Arial", 9),
            bg="#ecf0f1",
            fg="#2c3e50",
            anchor="w",
            padx=10,
            pady=5
        )
        self.status_label.pack(side=tk.BOTTOM, fill=tk.X)
        
    def browse_file(self):
        """Ch?n file PowerPoint"""
        filename = filedialog.askopenfilename(
            title="Ch?n file PowerPoint",
            filetypes=[
                ("PowerPoint files", "*.pptx"),
                ("All files", "*.*")
            ]
        )
        if filename:
            self.file_path.set(filename)
            self.update_status(f"?? ch?n: {os.path.basename(filename)}")
    
    def update_status(self, message):
        """C?p nh?t status bar"""
        self.status_label.config(text=message)
        self.root.update()
    
    def check_image_match(self, width_inches, height_inches):
        """Ki?m tra xem h?nh ?nh c? kh?p v?i ?i?u ki?n kh?ng"""
        target_width = self.width_value.get()
        target_height = self.height_value.get()
        tolerance = self.tolerance.get()
        mode = self.match_mode.get()
        
        width_match = abs(width_inches - target_width) < tolerance
        height_match = abs(height_inches - target_height) < tolerance
        
        if mode == "and":
            return width_match and height_match
        elif mode == "or":
            return width_match or height_match
        elif mode == "width_only":
            return width_match
        elif mode == "height_only":
            return height_match
        
        return False
    
    def preview_images(self):
        """Xem tr??c c?c h?nh s? b? x?a"""
        if not self.file_path.get():
            messagebox.showwarning("C?nh b?o", "Vui l?ng ch?n file PowerPoint!")
            return
        
        try:
            self.update_status("?ang ph?n t?ch file...")
            prs = Presentation(self.file_path.get())
            
            images_to_delete = []
            total_images = 0
            
            for slide_idx, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.shape_type == 13:  # Picture type
                        total_images += 1
                        width_inches = shape.width / 914400  # EMUs to inches
                        height_inches = shape.height / 914400
                        
                        if self.check_image_match(width_inches, height_inches):
                            images_to_delete.append({
                                'slide': slide_idx,
                                'name': shape.name,
                                'width': width_inches,
                                'height': height_inches
                            })
            
            # Show preview window
            self.show_preview_window(images_to_delete, total_images)
            self.update_status("S?n s?ng")
            
        except Exception as e:
            messagebox.showerror("L?i", f"Kh?ng th? ??c file:\n{str(e)}")
            self.update_status("L?i!")
    
    def show_preview_window(self, images_to_delete, total_images):
        """Hi?n th? c?a s? preview"""
        preview_window = tk.Toplevel(self.root)
        preview_window.title("Preview - Xem tr??c")
        preview_window.geometry("600x500")
        
        # Header
        header = tk.Label(
            preview_window,
            text=f"T?m th?y {len(images_to_delete)} h?nh ?nh s? b? x?a (T?ng: {total_images})",
            font=("Arial", 12, "bold"),
            bg="#e74c3c",
            fg="white",
            pady=10
        )
        header.pack(fill=tk.X)
        
        # Listbox with scrollbar
        list_frame = tk.Frame(preview_window)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        scrollbar = tk.Scrollbar(list_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        listbox = tk.Listbox(
            list_frame,
            font=("Courier", 9),
            yscrollcommand=scrollbar.set
        )
        listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.config(command=listbox.yview)
        
        # Add items
        if images_to_delete:
            for img in images_to_delete:
                item = f"Slide {img['slide']:3d} | {img['name']:30s} | {img['width']:.2f}x{img['height']:.2f} inch"
                listbox.insert(tk.END, item)
        else:
            listbox.insert(tk.END, "Kh?ng t?m th?y h?nh ?nh n?o ph? h?p!")
        
        # Close button
        close_btn = tk.Button(
            preview_window,
            text="??ng",
            command=preview_window.destroy,
            bg="#95a5a6",
            fg="white",
            font=("Arial", 10),
            padx=30,
            pady=5
        )
        close_btn.pack(pady=10)
    
    def process_file(self):
        """X? l? x?a h?nh ?nh"""
        if not self.file_path.get():
            messagebox.showwarning("C?nh b?o", "Vui l?ng ch?n file PowerPoint!")
            return
        
        # Confirm
        mode_text = {
            "and": "C? width V? height kh?p",
            "or": "width HO?C height kh?p",
            "width_only": "ch? width kh?p",
            "height_only": "ch? height kh?p"
        }
        
        confirm_msg = f"""B?n c? ch?c ch?n mu?n x?a h?nh ?nh v?i c?i ??t:

Width: {self.width_value.get()} inch
Height: {self.height_value.get()} inch
Ch? ??: {mode_text[self.match_mode.get()]}
Dung sai: {self.tolerance.get()} inch

Thao t?c n?y kh?ng th? ho?n t?c!
(File backup s? ???c t?o n?u b?n b?t t? ??ng backup)"""
        
        if not messagebox.askyesno("X?c nh?n", confirm_msg):
            return
        
        try:
            file_path = self.file_path.get()
            
            # Backup if enabled
            if self.auto_backup.get():
                self.update_status("?ang backup file g?c...")
                backup_path = self.create_backup(file_path)
                self.update_status(f"?? backup: {os.path.basename(backup_path)}")
            
            # Process
            self.update_status("?ang x? l? file...")
            prs = Presentation(file_path)
            
            deleted_count = 0
            total_images = 0
            
            for slide in prs.slides:
                shapes_to_delete = []
                
                for shape in slide.shapes:
                    if shape.shape_type == 13:  # Picture type
                        total_images += 1
                        width_inches = shape.width / 914400
                        height_inches = shape.height / 914400
                        
                        if self.check_image_match(width_inches, height_inches):
                            shapes_to_delete.append(shape)
                
                # Delete shapes
                for shape in shapes_to_delete:
                    sp = shape.element
                    sp.getparent().remove(sp)
                    deleted_count += 1
            
            # Save file
            self.update_status("?ang l?u file...")
            prs.save(file_path)
            
            # Show result
            messagebox.showinfo(
                "Ho?n t?t!",
                f"?? x?a th?nh c?ng!\n\n"
                f"T?ng s? h?nh ?nh: {total_images}\n"
                f"?? x?a: {deleted_count} h?nh ?nh\n"
                f"C?n l?i: {total_images - deleted_count} h?nh ?nh"
            )
            
            self.update_status("Ho?n t?t!")
            
        except Exception as e:
            messagebox.showerror("L?i", f"Kh?ng th? x? l? file:\n{str(e)}")
            self.update_status("L?i!")
    
    def create_backup(self, file_path):
        """T?o file backup"""
        file_dir = os.path.dirname(file_path)
        file_name = os.path.basename(file_path)
        name_without_ext = os.path.splitext(file_name)[0]
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"{name_without_ext}_backup_{timestamp}.pptx"
        backup_path = os.path.join(file_dir, backup_name)
        
        shutil.copy2(file_path, backup_path)
        return backup_path


def main():
    root = tk.Tk()
    app = PowerPointImageRemover(root)
    root.mainloop()


if __name__ == "__main__":
    main()
